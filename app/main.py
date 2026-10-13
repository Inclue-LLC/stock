import io
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path
from urllib.parse import urlencode

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from . import importers, services, xlsx_templates
from . import jan as jancode
from .db import get_db, init_db
from .models import (
    ImportBatch,
    Location,
    LocationType,
    Movement,
    MovementKind,
    Product,
    Region,
    Source,
    Stock,
)
from .services import StockError

APP_DIR = Path(__file__).resolve().parent
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


@asynccontextmanager
async def lifespan(_app):
    init_db()
    yield


app = FastAPI(title="在庫管理", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")
templates = Jinja2Templates(directory=APP_DIR / "templates")
templates.env.globals.update(Region=Region, LocationType=LocationType, MovementKind=MovementKind)


def render(request: Request, name: str, **ctx):
    def url_with(**changes):
        params = {k: v for k, v in request.query_params.items() if k not in ("page", "msg", "error")}
        params.update(changes)
        params = {k: v for k, v in params.items() if v not in (None, "")}
        return "?" + urlencode(params) if params else request.url.path

    ctx.setdefault("url_with", url_with)
    ctx.setdefault("path", request.url.path)
    ctx.setdefault("msg", request.query_params.get("msg"))
    ctx.setdefault("error", request.query_params.get("error"))
    return templates.TemplateResponse(request, name, ctx)


def redirect(path: str, **params) -> RedirectResponse:
    params = {k: v for k, v in params.items() if v}
    return RedirectResponse(f"{path}?{urlencode(params)}" if params else path, status_code=303)


def enum_or_none(enum_cls, value: str | None):
    try:
        return enum_cls(value) if value else None
    except ValueError:
        return None


def int_or_none(value: str | None) -> int | None:
    try:
        return int(value) if value else None
    except ValueError:
        return None


def parse_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def page_param(request: Request) -> int:
    return int_or_none(request.query_params.get("page")) or 1


def xlsx_response(wb, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def get_or_404(db: Session, model, obj_id: int):
    obj = db.get(model, obj_id)
    if obj is None:
        raise HTTPException(404)
    return obj


def all_locations(db: Session) -> list[Location]:
    return list(db.scalars(select(Location).order_by(Location.region, Location.type, Location.name)))


@app.get("/")
def home():
    return RedirectResponse("/stocks", status_code=307)


# ---- 在庫一覧 -------------------------------------------------------------------

def _stock_filters(request: Request):
    p = request.query_params
    return dict(
        region=enum_or_none(Region, p.get("region")),
        ltype=enum_or_none(LocationType, p.get("type")),
        location_id=int_or_none(p.get("location_id")),
        q=(p.get("q") or "").strip(),
        include_zero=p.get("include_zero") == "1",
    )


@app.get("/stocks")
def stock_list(request: Request, db: Session = Depends(get_db)):
    filters = _stock_filters(request)
    view = "location" if request.query_params.get("view") == "location" else "product"
    if view == "product":
        page = services.paginate(
            db, services.product_totals_query(filters["q"], only_in_stock=not filters["include_zero"]),
            page_param(request), per_page=100, scalars=False,
        )
    else:
        page = services.paginate(db, services.stock_query(**filters), page_param(request), per_page=100)
    return render(
        request, "stocks.html",
        view=view, page=page, filters=filters, locations=all_locations(db),
        categories=services.CATEGORIES,
        export_query=request.url.query,
    )


@app.get("/stocks/export.xlsx")
def stock_export(request: Request, db: Session = Depends(get_db)):
    from openpyxl import Workbook

    filters = _stock_filters(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "在庫"
    if request.query_params.get("view") == "location":
        ws.append(["JANコード", "品番", "商品名", "地域", "種別", "拠点", "数量", "単位"])
        for s in db.scalars(services.stock_query(**filters)):
            ws.append([s.product.jan, s.product.sku, s.product.name, s.location.region.label,
                       s.location.type.label, s.location.name, s.quantity, s.product.unit])
    else:
        ws.append(["JANコード", "品番", "商品名"]
                  + [f"{r.label}・{t.label}" for r, t in services.CATEGORIES] + ["合計", "単位"])
        for p, *qtys in db.execute(services.product_totals_query(filters["q"], not filters["include_zero"])):
            ws.append([p.jan, p.sku, p.name, *qtys, p.unit])
    for col in ws.iter_cols(min_col=1, max_col=1):
        for c in col:
            c.number_format = "@"  # JAN を文字列として扱う
    return xlsx_response(wb, "stock.xlsx")


# ---- 拠点 ----------------------------------------------------------------------

@app.get("/locations")
def location_list(request: Request, db: Session = Depends(get_db)):
    last_updated = dict(db.execute(
        select(Movement.location_id, func.max(Movement.created_at)).group_by(Movement.location_id)
    ).all())
    totals = dict(db.execute(
        select(Stock.location_id, func.sum(Stock.quantity)).group_by(Stock.location_id)
    ).all())
    return render(request, "locations.html", locations=all_locations(db), last_updated=last_updated,
                  totals=totals)


def _store_report(request: Request, db: Session, location: Location):
    p = request.query_params
    return services.store_report(db, location, brand=p.get("brand", ""), q=(p.get("q") or "").strip())


@app.get("/locations/{location_id}")
def location_detail(request: Request, location_id: int, db: Session = Depends(get_db)):
    location = get_or_404(db, Location, location_id)
    report = _store_report(request, db, location)
    last = db.scalar(select(func.max(Movement.created_at)).where(Movement.location_id == location.id))
    return render(request, "location_detail.html", location=location, report=report, last_updated=last,
                  brand=request.query_params.get("brand", ""), q=request.query_params.get("q", ""),
                  show_dates=request.query_params.get("dates") != "0")


@app.get("/locations/{location_id}/export.xlsx")
def location_export(request: Request, location_id: int, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    location = get_or_404(db, Location, location_id)
    report = _store_report(request, db, location)
    wb = Workbook()
    ws = wb.active
    ws.title = location.name[:31]
    dates = report.delivery_dates
    ws.append(["ブランド", "商品名", "JANコード", "上代（税込）"]
              + [f"{d:%y/%m/%d} 納品" for d in dates]
              + ["納品数 合計", "販売数", "在庫更新差分", "在庫数", "消化率", "販売金額（上代）", "在庫金額（上代）"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E9E")
    for r in report.rows:
        ws.append([r.product.brand, r.product.name, r.product.jan, r.product.price]
                  + [r.deliveries.get(d) for d in dates]
                  + [r.delivered, r.sold, r.adjusted or None, r.stock, r.sell_through,
                     r.sales_amount, r.stock_amount])
    last_col = ws.max_column
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "@"
        row[last_col - 3].number_format = "0.0%"
        for c in (row[3], row[last_col - 2], row[last_col - 1]):
            c.number_format = "#,##0"
    ws.freeze_panes = "C2"
    for col, width in zip("ABCD", (18, 40, 16, 12)):
        ws.column_dimensions[col].width = width
    return xlsx_response(wb, "store_report.xlsx")


@app.post("/locations")
def location_create(
    name: str = Form(...), region: Region = Form(...), type: LocationType = Form(...),
    note: str = Form(""), db: Session = Depends(get_db),
):
    name = name.strip()
    if not name:
        return redirect("/locations", error="拠点名を入力してください")
    db.add(Location(name=name, region=region, type=type, note=note))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect("/locations", error=f"拠点「{name}」はすでに登録されています")
    return redirect("/locations", msg=f"拠点「{name}」を登録しました")


@app.post("/locations/{location_id}")
def location_update(
    location_id: int, name: str = Form(...), region: Region = Form(...), type: LocationType = Form(...),
    note: str = Form(""), db: Session = Depends(get_db),
):
    loc = get_or_404(db, Location, location_id)
    loc.name, loc.region, loc.type, loc.note = name.strip(), region, type, note
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect("/locations", error=f"拠点「{name}」はすでに登録されています")
    return redirect("/locations", msg="拠点を更新しました")


@app.post("/locations/{location_id}/delete")
def location_delete(location_id: int, db: Session = Depends(get_db)):
    loc = get_or_404(db, Location, location_id)
    if db.scalar(select(func.count(Movement.id)).where(Movement.location_id == loc.id)):
        return redirect("/locations", error="入出庫履歴がある拠点は削除できません")
    db.delete(loc)
    db.commit()
    return redirect("/locations", msg="拠点を削除しました")


# ---- 商品 ----------------------------------------------------------------------

@app.get("/products")
def product_list(request: Request, q: str = "", db: Session = Depends(get_db)):
    stmt = services.product_search(
        select(Product, func.coalesce(func.sum(Stock.quantity), 0))
        .outerjoin(Stock, Stock.product_id == Product.id)
        .group_by(Product.id)
        .order_by(Product.name, Product.jan),
        q,
    )
    page = services.paginate(db, stmt, page_param(request), per_page=50, scalars=False)
    brands = db.scalars(select(Product.brand).where(Product.brand != "").distinct().order_by(Product.brand)).all()
    return render(request, "products.html", page=page, q=q, brands=brands,
                  prefill_jan=request.query_params.get("jan", ""))


def price_or_none(value: str) -> int | None:
    return importers.parse_quantity(value) if value.strip() else None


@app.post("/products")
def product_create(
    jan: str = Form(...), name: str = Form(...), sku: str = Form(""), unit: str = Form("個"),
    brand: str = Form(""), price: str = Form(""), note: str = Form(""), db: Session = Depends(get_db),
):
    code, err = jancode.try_normalize(jan)
    if err:
        return redirect("/products", error=err, jan=jan)
    name = name.strip()
    if not name:
        return redirect("/products", error="商品名を入力してください", jan=jan)
    product = Product(jan=code, name=name, sku=sku.strip(), unit=unit.strip() or "個", note=note,
                      brand=brand.strip(), price=price_or_none(price))
    db.add(product)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect("/products", error=f"JANコード「{code}」はすでに登録されています")
    return redirect("/products", msg=f"商品「{name}」を登録しました")


@app.post("/products/import")
async def product_import(file: UploadFile = File(...), db: Session = Depends(get_db)):
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        return redirect("/products", error="ファイルサイズが大きすぎます（上限20MB）")
    try:
        parsed = importers.parse_product_file(file.filename or "", data)
    except ValueError as e:
        return redirect("/products", error=str(e))
    except Exception:
        return redirect("/products", error="ファイルを読み込めませんでした。形式を確認してください")
    if not parsed.rows:
        return redirect("/products", error=" / ".join(parsed.warnings) or "商品が見つかりませんでした")
    result = services.import_products(db, parsed.rows)
    db.commit()
    msg = f"商品マスタを取り込みました：新規 {result.created}件、更新 {result.updated}件"
    error = None
    if result.errors:
        shown = result.errors[:10]
        more = f" ほか{len(result.errors) - 10}件" if len(result.errors) > 10 else ""
        error = f"エラー {len(result.errors)}件（取り込まれていません）: " + " / ".join(shown) + more
    return redirect("/products", msg=msg, error=error)


@app.get("/products/template.xlsx")
def product_template():
    return xlsx_response(xlsx_templates.product_master_workbook(), "product_master.xlsx")


@app.get("/products/{product_id}")
def product_detail(request: Request, product_id: int, db: Session = Depends(get_db)):
    product = get_or_404(db, Product, product_id)
    stocks = db.scalars(
        select(Stock).join(Stock.location).where(Stock.product_id == product.id)
        .order_by(Location.region, Location.type, Location.name)
    ).all()
    movements = db.scalars(
        select(Movement).where(Movement.product_id == product.id)
        .order_by(func.coalesce(Movement.date, func.date(Movement.created_at)).desc(), Movement.id.desc()).limit(100)
    ).all()
    return render(request, "product_detail.html", product=product, stocks=stocks, movements=movements,
                  total=sum(s.quantity for s in stocks))


@app.post("/products/{product_id}")
def product_update(
    product_id: int, jan: str = Form(...), name: str = Form(...), sku: str = Form(""), unit: str = Form("個"),
    brand: str = Form(""), price: str = Form(""), note: str = Form(""), db: Session = Depends(get_db),
):
    product = get_or_404(db, Product, product_id)
    code, err = jancode.try_normalize(jan)
    if err:
        return redirect(f"/products/{product_id}", error=err)
    if not name.strip():
        return redirect(f"/products/{product_id}", error="商品名を入力してください")
    product.jan, product.sku, product.name = code, sku.strip(), name.strip()
    product.unit, product.note = unit.strip() or "個", note
    product.brand, product.price = brand.strip(), price_or_none(price)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect(f"/products/{product_id}", error=f"JANコード「{code}」はすでに使われています")
    return redirect(f"/products/{product_id}", msg="商品を更新しました")


@app.post("/products/{product_id}/delete")
def product_delete(product_id: int, db: Session = Depends(get_db)):
    product = get_or_404(db, Product, product_id)
    if db.scalar(select(func.count(Movement.id)).where(Movement.product_id == product.id)):
        return redirect(f"/products/{product_id}", error="入出庫履歴がある商品は削除できません")
    db.delete(product)
    db.commit()
    return redirect("/products", msg="商品を削除しました")


# ---- 入出庫履歴 -------------------------------------------------------------------

@app.get("/movements")
def movement_list(request: Request, db: Session = Depends(get_db)):
    p = request.query_params
    stmt = select(Movement).join(Movement.location).join(Movement.product).order_by(Movement.id.desc())
    region = enum_or_none(Region, p.get("region"))
    ltype = enum_or_none(LocationType, p.get("type"))
    location_id = int_or_none(p.get("location_id"))
    source = enum_or_none(Source, p.get("source"))
    batch_id = int_or_none(p.get("batch_id"))
    q = (p.get("q") or "").strip()
    if region:
        stmt = stmt.where(Location.region == region)
    if ltype:
        stmt = stmt.where(Location.type == ltype)
    if location_id:
        stmt = stmt.where(Location.id == location_id)
    if source:
        stmt = stmt.where(Movement.source == source)
    if batch_id:
        stmt = stmt.where(Movement.batch_id == batch_id)
    stmt = services.product_search(stmt, q)
    page = services.paginate(db, stmt, page_param(request), per_page=100)
    return render(request, "movements.html", page=page, locations=all_locations(db),
                  filters=dict(region=region, ltype=ltype, location_id=location_id, source=source, q=q,
                               batch_id=batch_id),
                  Source=Source)


# ---- ファイル取込（Excel / CSV / PDF） ---------------------------------------------

IMPORT_KINDS = (MovementKind.SET, MovementKind.IN, MovementKind.OUT)


@app.get("/imports")
def import_index(request: Request, db: Session = Depends(get_db)):
    batches = db.scalars(select(ImportBatch).order_by(ImportBatch.id.desc()).limit(50)).all()
    return render(request, "imports.html", batches=batches, locations=all_locations(db), kinds=IMPORT_KINDS,
                  selected_location=int_or_none(request.query_params.get("location_id")),
                  selected_kind=request.query_params.get("kind") or "SET", today=date.today().isoformat())


@app.get("/imports/template.xlsx")
def import_template():
    return xlsx_response(xlsx_templates.stock_import_workbook(), "stock_import_template.xlsx")


@app.post("/imports")
async def import_upload(
    file: UploadFile = File(...),
    kind: str = Form(...),
    location_id: str = Form(""),
    zero_missing: str = Form(""),
    on: str = Form(""),
    db: Session = Depends(get_db),
):
    mk = enum_or_none(MovementKind, kind)
    if mk not in IMPORT_KINDS:
        return redirect("/imports", error="区分が不正です")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        return redirect("/imports", error="ファイルサイズが大きすぎます（上限20MB）")
    filename = file.filename or "upload"
    try:
        source, parsed = importers.parse_file(filename, data)
    except ValueError as e:
        return redirect("/imports", error=str(e))
    except Exception:
        return redirect("/imports", error="ファイルを読み込めませんでした。形式を確認してください")
    loc_id = int_or_none(location_id)
    if loc_id is not None:
        get_or_404(db, Location, loc_id)
    batch = ImportBatch(filename=filename, source=Source(source), kind=mk, location_id=loc_id,
                        rows=parsed.rows, warnings=parsed.warnings,
                        zero_missing=(mk == MovementKind.SET and zero_missing == "1"),
                        date=parse_date(on) or date.today())
    db.add(batch)
    db.commit()
    return redirect(f"/imports/{batch.id}")


@app.get("/imports/{batch_id}")
def import_preview(request: Request, batch_id: int, db: Session = Depends(get_db)):
    batch = get_or_404(db, ImportBatch, batch_id)
    preview = services.validate_batch(db, batch) if batch.status == "pending" else []
    missing = services.missing_stocks(db, batch, preview) if preview else []
    return render(
        request, "import_preview.html", batch=batch, preview=preview, locations=all_locations(db),
        missing=missing,
        ok_count=sum(1 for p in preview if not p["errors"]),
        new_products=sum(1 for p in preview if not p["errors"] and p["product"] is None),
    )


@app.post("/imports/{batch_id}/location")
def import_set_location(batch_id: int, location_id: str = Form(""), db: Session = Depends(get_db)):
    batch = get_or_404(db, ImportBatch, batch_id)
    if batch.status == "pending":
        batch.location_id = int_or_none(location_id)
        db.commit()
    return redirect(f"/imports/{batch_id}")


@app.post("/imports/{batch_id}/apply")
def import_apply(batch_id: int, db: Session = Depends(get_db)):
    batch = get_or_404(db, ImportBatch, batch_id)
    try:
        applied, skipped, zeroed = services.apply_batch(db, batch)
        db.commit()
    except StockError as e:
        db.rollback()
        return redirect(f"/imports/{batch_id}", error=str(e))
    msg = f"{applied}件を反映しました" + (f"（エラー {skipped}件はスキップ）" if skipped else "")
    if zeroed:
        msg += f"。ファイルにない {zeroed}件を0にしました"
    return redirect(f"/imports/{batch_id}", msg=msg)


@app.post("/imports/{batch_id}/cancel")
def import_cancel(batch_id: int, db: Session = Depends(get_db)):
    batch = get_or_404(db, ImportBatch, batch_id)
    if batch.status == "pending":
        batch.status = "cancelled"
        db.commit()
    return redirect("/imports", msg="取込を取り消しました")
