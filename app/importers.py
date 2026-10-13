"""Excel / CSV / PDF から在庫明細（SKU・商品名・数量・拠点）を読み取る。

ヘッダー行は先頭付近から自動検出する。日本語・中国語・英語の列名に対応。
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

HEADER_ALIASES: dict[str, list[str]] = {
    "sku": ["sku", "品番", "商品コード", "商品番号", "型番", "品目コード", "コード", "jan", "janコード",
            "item code", "itemcode", "code", "货号", "商品编码", "编码", "编号", "产品编号"],
    "name": ["商品名", "品名", "品目", "品目名", "名称", "製品名", "item", "name", "product", "description",
             "商品名称", "产品名称", "品名称"],
    "quantity": ["数量", "個数", "在庫数", "在庫", "入庫数", "出庫数", "販売数", "売上数", "qty", "quantity",
                 "count", "数", "库存", "库存数量", "入库数量", "出库数量", "销售数量"],
    "location": ["拠点", "倉庫", "店舗", "保管場所", "委託先", "location", "warehouse", "store", "仓库", "门店", "店铺"],
    "unit": ["単位", "unit", "单位"],
    "note": ["備考", "メモ", "note", "remarks", "备注"],
}

REQUIRED = ("sku", "quantity")
HEADER_SCAN_ROWS = 30


def normalize(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(text))).lower()


_ALIAS_LOOKUP = {normalize(a): key for key, aliases in HEADER_ALIASES.items() for a in aliases}


def parse_quantity(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else None
    s = normalize(value).replace(",", "")
    m = re.fullmatch(r"(-?\d+(?:\.0+)?)[^\d]*", s)
    if not m:
        return None
    return int(float(m.group(1)))


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def match_header(cells: list) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(cells):
        key = _ALIAS_LOOKUP.get(normalize(cell))
        if key and key not in mapping:
            mapping[key] = idx
    if all(k in mapping for k in REQUIRED):
        return mapping
    return None


def rows_from_table(table: list[list], result: ParseResult, label: str = "") -> bool:
    """2次元配列からヘッダーを検出して明細行を抽出する。ヘッダーが見つかれば True。"""
    header_idx = mapping = None
    for i, cells in enumerate(table[:HEADER_SCAN_ROWS]):
        mapping = match_header(list(cells))
        if mapping:
            header_idx = i
            break
    if mapping is None:
        return False

    def cell(cells, key):
        idx = mapping.get(key)
        if idx is None or idx >= len(cells):
            return None
        v = cells[idx]
        return v.strip() if isinstance(v, str) else v

    for offset, cells in enumerate(table[header_idx + 1:], start=header_idx + 2):
        cells = list(cells)
        sku = cell(cells, "sku")
        if sku is None or str(sku).strip() == "":
            continue
        sku = unicodedata.normalize("NFKC", str(sku)).strip()
        if isinstance(cell(cells, "sku"), float) and float(sku).is_integer():
            sku = str(int(float(sku)))
        if normalize(sku) in ("合計", "小計", "total", "合计", "小计"):
            continue
        raw_qty = cell(cells, "quantity")
        qty = parse_quantity(raw_qty)
        if qty is None:
            result.warnings.append(f"{label}{offset}行目: 数量「{raw_qty}」を読み取れません")
        name = cell(cells, "name")
        location = cell(cells, "location")
        result.rows.append({
            "sku": sku,
            "name": str(name) if name not in (None, "") else "",
            "quantity": qty,
            "location": str(location) if location not in (None, "") else None,
            "unit": str(cell(cells, "unit") or ""),
            "note": str(cell(cells, "note") or ""),
        })
    return True


def parse_excel(data: bytes) -> ParseResult:
    from openpyxl import load_workbook

    result = ParseResult()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    found = False
    for ws in wb.worksheets:
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        label = f"[{ws.title}] " if len(wb.worksheets) > 1 else ""
        if rows_from_table(table, result, label):
            found = True
    if not found:
        result.warnings.append("ヘッダー行（品番・数量 など）が見つかりませんでした")
    return result


def parse_csv(data: bytes) -> ParseResult:
    for enc in ("utf-8-sig", "cp932", "gb18030"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return ParseResult(warnings=["文字コードを判別できませんでした"])
    result = ParseResult()
    table = list(csv.reader(io.StringIO(text)))
    if not rows_from_table(table, result):
        result.warnings.append("ヘッダー行（品番・数量 など）が見つかりませんでした")
    return result


# 表として抽出できない PDF 用: 「品番 商品名 数量」の並びの行を拾う
_LINE_RE = re.compile(r"^(?P<sku>[A-Za-z0-9][\w\-./]*)\s+(?P<name>.+?)\s+(?P<qty>-?\d[\d,]*)\s*\S{0,3}$")


def parse_pdf(data: bytes) -> ParseResult:
    import pdfplumber

    result = ParseResult()
    found = False
    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                if rows_from_table(table, result, f"p{pno} "):
                    found = True
            lines.extend((page.extract_text() or "").splitlines())
    if found:
        return result

    for line in lines:
        m = _LINE_RE.match(unicodedata.normalize("NFKC", line).strip())
        if m:
            result.rows.append({
                "sku": m["sku"], "name": m["name"].strip(), "quantity": parse_quantity(m["qty"]),
                "location": None, "unit": "", "note": "",
            })
    if result.rows:
        result.warnings.append("PDF内に表が見つからなかったため、テキスト行から推定しました。内容を必ず確認してください")
    else:
        result.warnings.append("PDFから明細を読み取れませんでした（スキャン画像のPDFは非対応です）")
    return result


def parse_file(filename: str, data: bytes) -> tuple[str, ParseResult]:
    """戻り値: (Source の値, ParseResult)"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("xlsx", "xlsm"):
        return "EXCEL", parse_excel(data)
    if ext == "csv":
        return "CSV", parse_csv(data)
    if ext == "pdf":
        return "PDF", parse_pdf(data)
    raise ValueError("対応形式は .xlsx / .xlsm / .csv / .pdf です")
